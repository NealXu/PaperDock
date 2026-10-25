from __future__ import annotations

import shutil
import time
from pathlib import Path

import pymupdf
import pypdfium2 as pdfium
import pytest
from docx import Document
from pptx import Presentation

from conftest import OCR_PACK, SAMPLE
from helpers import AppHelpers


def _new_files(output: Path, before: set[str]):
    return sorted(path for path in output.glob("*.pdf") if path.name not in before)


def _pdfium_text(path: Path, page_index: int = 0) -> str:
    document = pdfium.PdfDocument(str(path))
    try:
        return document[page_index].get_textpage().get_text_range()
    finally:
        document.close()


def test_split_modes_searchable_and_watermark(session, app_helpers):
    win = session.window
    helpers = app_helpers
    before = set()

    helpers.finish_operation(win, "StartSplit")
    fixed_files = _new_files(session.output, before)
    assert len(fixed_files) == 3
    assert all(pymupdf.open(path).page_count == 1 for path in fixed_files)
    before = {path.name for path in session.output.glob("*.pdf")}

    helpers.select_radio(win, "SplitMode1")
    helpers.set_text(win, "SplitRanges", "1;2-3")
    helpers.finish_operation(win, "StartSplit")
    custom_files = _new_files(session.output, before)
    assert len(custom_files) == 2
    assert sorted(pymupdf.open(path).page_count for path in custom_files) == [1, 2]
    before = {path.name for path in session.output.glob("*.pdf")}

    helpers.select_radio(win, "SplitMode2")
    helpers.set_text(win, "SplitPartCount", "2")
    helpers.select_radio(win, "SplitOutputImage")
    helpers.finish_operation(win, "StartSplit")
    count_files = _new_files(session.output, before)
    assert len(count_files) == 2
    assert all(pymupdf.open(path).page_count in (1, 2) for path in count_files)
    assert all(not pymupdf.open(path)[0].get_text().strip() for path in count_files)
    before = {path.name for path in session.output.glob("*.pdf")}

    helpers.select_radio(win, "SplitOutputSearchable")
    helpers.finish_operation(win, "StartSplit", timeout=300)
    searchable_files = _new_files(session.output, before)
    assert len(searchable_files) == 2
    assert all(Path(path).stat().st_size > 50_000 for path in searchable_files)
    assert any("SEARCHABLE" in _pdfium_text(path) for path in searchable_files)
    before = {path.name for path in session.output.glob("*.pdf")}

    more = win.child_window(title="更多选项", control_type="Button")
    more.click_input()
    watermark = win.child_window(auto_id="SplitWatermarkEnabled", control_type="CheckBox")
    watermark.wait("visible", timeout=10)
    watermark.click_input()
    helpers.set_text(win, "SplitWatermarkText", "E2E-WATERMARK")
    helpers.select_radio(win, "SplitOutputPdf")
    helpers.finish_operation(win, "StartSplit")
    watermarked = _new_files(session.output, before)
    assert len(watermarked) == 2
    assert any("E2E-WATERMARK" in pymupdf.open(path)[0].get_text() for path in watermarked)
    helpers.screenshot(win, session.shots / "split-complete.png")


def test_merge_with_per_file_ranges(session, app_helpers):
    win = session.window
    helpers = app_helpers
    second = session.root / "附录 与 补充.pdf"
    shutil.copy2(SAMPLE, second)

    win.child_window(auto_id="Navmerge", control_type="Button").invoke()
    helpers.choose_file(win, "AddMergeFiles", "选择要合并的 PDF", SAMPLE)
    win.child_window(auto_id="MergeRange0", control_type="Edit").wait("visible", timeout=20)
    helpers.choose_file(win, "AddMergeFiles", "选择要合并的 PDF", second)
    win.child_window(auto_id="MergeRange1", control_type="Edit").wait("visible", timeout=20)

    helpers.set_text(win, "MergeRange0", "1-2")
    helpers.set_text(win, "MergeRange1", "2")
    helpers.set_text(win, "MergeOutputName", "merged-e2e.pdf")
    win.child_window(auto_id="MergeOutputName", control_type="Edit").click_input()
    helpers.finish_operation(win, "StartMerge")

    merged = session.output / "merged-e2e.pdf"
    assert merged.exists()
    assert pymupdf.open(merged).page_count == 3
    assert "PAGE 2" in pymupdf.open(merged)[2].get_text()
    helpers.screenshot(win, session.shots / "merge-complete.png")


def test_convert_to_word_and_powerpoint(session, app_helpers):
    win = session.window
    helpers = app_helpers
    win.child_window(auto_id="Navconvert", control_type="Button").invoke()
    helpers.wait_status(win, "就绪 · 3 页", 30)

    targets = [
        ("ConvertWord", "convert-e2e.docx"),
        ("ConvertPowerPoint", "convert-e2e.pptx"),
    ]
    for target_id, name in targets:
        helpers.select_radio(win, target_id)
        if target_id == "ConvertWord":
            helpers.select_radio(win, "WordModeEditable")
        output = session.output / name
        helpers.set_text(win, "ConvertOutputPath", str(output))
        helpers.finish_operation(win, "StartConvert", timeout=240)
        assert output.exists()

    text = "\n".join(paragraph.text for paragraph in Document(session.output / "convert-e2e.docx").paragraphs)
    assert "PAGE 1" in text
    assert len(Presentation(session.output / "convert-e2e.pptx").slides) == 3
    helpers.screenshot(win, session.shots / "convert-complete.png")


def test_word_fidelity_preserves_a4_pages(session, app_helpers):
    win = session.window
    helpers = app_helpers
    win.child_window(auto_id="Navconvert", control_type="Button").invoke()
    helpers.wait_status(win, "就绪 · 3 页", 30)

    helpers.select_radio(win, "ConvertWord")
    fidelity = win.child_window(auto_id="WordModeFidelity", control_type="RadioButton")
    fidelity.wait("visible", timeout=10)
    fidelity.click_input()
    output = session.output / "word-fidelity.docx"
    helpers.set_text(win, "ConvertOutputPath", str(output))
    helpers.finish_operation(win, "StartConvert", timeout=240)

    document = Document(output)
    section = document.sections[0]
    assert round(section.page_width.pt) == 595
    assert round(section.page_height.pt) == 842
    assert len(document.inline_shapes) == 3
    assert "第 1 页" not in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_editable_word_uses_source_page_size_without_injected_labels(session, app_helpers):
    win = session.window
    helpers = app_helpers
    win.child_window(auto_id="Navconvert", control_type="Button").invoke()
    helpers.wait_status(win, "就绪 · 3 页", 30)

    helpers.select_radio(win, "ConvertWord")
    helpers.select_radio(win, "WordModeEditable")
    output = session.output / "word-editable.docx"
    helpers.set_text(win, "ConvertOutputPath", str(output))
    helpers.finish_operation(win, "StartConvert", timeout=240)

    document = Document(output)
    section = document.sections[0]
    assert round(section.page_width.pt) == 595
    assert round(section.page_height.pt) == 842
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "PAGE 1" in text
    assert "第 1 页" not in text


@pytest.mark.skipif(not OCR_PACK.exists(), reason="offline OCR test pack not available")
def test_import_offline_ocr_pack_and_searchable_output(session, app_helpers):
    win = session.window
    helpers = app_helpers
    win.child_window(auto_id="Navsettings", control_type="Button").invoke()
    helpers.choose_file(win, "ImportOcrPack", "导入离线 OCR 包", OCR_PACK)
    helpers.wait_status(win, "OCR 包已导入", 30)
    choice = win.child_window(auto_id="DefaultOcrChoice", control_type="ComboBox")
    assert "Tesseract" in choice.selected_text()

    win.child_window(auto_id="Navsplit", control_type="Button").invoke()
    helpers.select_radio(win, "SplitOutputSearchable")
    helpers.finish_operation(win, "StartSplit", timeout=300)
    searchable = [path for path in session.output.glob("*.pdf") if path.stat().st_size > 50_000]
    assert len(searchable) == 3
    assert any("SEARCHABLE" in _pdfium_text(path) for path in searchable)


def test_rejects_invalid_ocr_pack(session, app_helpers, tmp_path):
    import zipfile

    win = session.window
    helpers = app_helpers
    invalid = tmp_path / "invalid.ocrpack.zip"
    with zipfile.ZipFile(invalid, "w") as archive:
        archive.writestr(
            "ocrpack.json",
            '{"id":"bad","name":"bad","version":"1","engine":"tesseract",'
            '"entry":"tesseract.exe","language":"chi_sim","dataDirectory":"tessdata"}',
        )
    win.child_window(auto_id="Navsettings", control_type="Button").invoke()
    helpers.choose_file(win, "ImportOcrPack", "导入离线 OCR 包", invalid)
    status = win.child_window(auto_id="StatusText", control_type="Text")
    deadline = time.time() + 15
    while time.time() < deadline and "失败" not in status.window_text():
        time.sleep(0.25)
    assert "失败" in status.window_text()
    time.sleep(0.5)
    helpers.dismiss_dialog(win)


def test_cancel_stops_long_searchable_job(session, app_helpers, tmp_path):
    win = session.window
    helpers = app_helpers

    source = pymupdf.open(SAMPLE)
    large = pymupdf.open()
    for _ in range(80):
        large.insert_pdf(source, from_page=0, to_page=0)
    large_path = tmp_path / "八十页扫描样本.pdf"
    large.save(large_path)
    large.close()
    source.close()

    helpers.choose_file(win, "HeaderAction", "选择 PDF", large_path)
    helpers.wait_status(win, "就绪 · 80 页", 30)
    helpers.select_radio(win, "SplitOutputSearchable")
    win.child_window(auto_id="StartSplit", control_type="Button").invoke()
    cancel = win.child_window(auto_id="CancelJob", control_type="Button")
    cancel.wait("visible", timeout=10)
    cancel.invoke()
    helpers.wait_status(win, "已取消", 30)
    assert not list(session.output.glob("*.pdf"))


def test_invalid_and_encrypted_pdf_fail_cleanly(session, app_helpers, tmp_path):
    win = session.window
    helpers = app_helpers
    status = win.child_window(auto_id="StatusText", control_type="Text")

    invalid = tmp_path / "损坏.pdf"
    invalid.write_bytes(b"not a pdf")
    helpers.choose_file(win, "HeaderAction", "选择 PDF", invalid)
    helpers.wait_status(win, "无法打开 PDF", 20)
    time.sleep(0.5)
    helpers.dismiss_dialog(win)

    encrypted = tmp_path / "加密.pdf"
    document = pymupdf.open(SAMPLE)
    document.save(
        encrypted,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-password",
        user_pw="user-password",
    )
    document.close()
    helpers.choose_file(win, "HeaderAction", "选择 PDF", encrypted)
    helpers.wait_status(win, "无法打开 PDF", 20)
    time.sleep(0.5)
    helpers.dismiss_dialog(win)
    assert "无法打开 PDF" in status.window_text()


def test_scanned_pdf_to_word_uses_ocr(session, app_helpers, tmp_path):
    source = pymupdf.open(SAMPLE)
    pixmap = source[0].get_pixmap(dpi=200, colorspace=pymupdf.csRGB, alpha=False)
    scanned = pymupdf.open()
    page = scanned.new_page(width=source[0].rect.width, height=source[0].rect.height)
    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
    scanned_path = tmp_path / "扫描页.pdf"
    scanned.save(scanned_path)
    scanned.close()
    source.close()

    win = session.window
    helpers = app_helpers
    helpers.choose_file(win, "HeaderAction", "选择 PDF", scanned_path)
    helpers.wait_status(win, "就绪 · 1 页", 30)
    win.child_window(auto_id="Navconvert", control_type="Button").invoke()
    helpers.wait_status(win, "就绪 · 1 页", 30)
    helpers.select_radio(win, "ConvertWord")
    helpers.select_radio(win, "WordModeEditable")
    output = session.output / "scanned-ocr.docx"
    helpers.set_text(win, "ConvertOutputPath", str(output))
    helpers.finish_operation(win, "StartConvert", timeout=240)
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "SEARCHABLE" in text
